import io
import uuid
from openpyxl import load_workbook
from sqlalchemy import select
from test_flows import app, TestClient, login, post, HEAD, SessionLocal
from apps.api.app.db import User, TicketVersion, Assignment, Attempt


def test_teacher_group_profile_stars_and_excel():
    with TestClient(app) as c:
        login(c,'admin')
        sid=post(c,'/users',{'login':'profile_'+uuid.uuid4().hex[:10],'name':'=Студент отчёта','password':'Training112!','role':'student'}).json()['id']
        post(c,'/logout');login(c,'teacher')
        group=post(c,'/groups',{'name':'Группа профилей'});assert group.status_code==200,group.text
        gid=group.json()['id']
        r=c.put(f'/api/teacher/groups/{gid}/members',json={'students':[sid]},headers=HEAD);assert r.status_code==200,r.text
        with SessionLocal() as db:
            teacher=db.scalar(select(User).where(User.login=='teacher'))
            version=db.scalar(select(TicketVersion).where(TicketVersion.published==True))
            assignment=Assignment(version_id=version.id,student_id=sid,teacher_id=teacher.id,title='Проверка профиля')
            db.add(assignment);db.flush()
            for i in range(4):
                db.add(Attempt(assignment_id=assignment.id,student_id=sid,task_index=i,status='completed',snapshot={**version.data['tasks'][0],'title':'Профиль','difficulty_level':3},card={},state={},assessment={'score':85 if i<3 else 0,'pending':i==3,'criteria':[{'skill':'Адрес','passed':False}]}))
            db.commit()
        row=next(x for x in c.get('/api/teacher/students').json() if x['id']==sid)
        assert row['score']==85 and row['confirmed']==3 and row['mastered_stars']==3
        assert row['errors'][0]['errors']==3 and row['can_assign']
        export=c.get(f'/api/teacher/report.xlsx?group_id={gid}');assert export.status_code==200
        pdf=c.get(f'/api/teacher/report.pdf?group_id={gid}');assert pdf.status_code==200 and pdf.content.startswith(b'%PDF-')
        ws=load_workbook(io.BytesIO(export.content)).active
        assert ws['A2'].value=='=Студент отчёта' and ws['A2'].data_type=='s'
        post(c,'/logout');login(c,'student')
        assert c.get('/api/teacher/students').status_code==403
        assert c.get('/api/teacher/report.xlsx').status_code==403
        assert post(c,'/groups',{'name':'Недоступно'}).status_code==403


def test_audio_and_demo_do_not_require_student_screen_sharing():
    with TestClient(app) as teacher,TestClient(app) as student:
        login(teacher,'admin')
        for client in (teacher,student):
            station=post(client,'/classroom/register',{'number':uuid.uuid4().hex[:10]}).json()
            client.headers['X-Workstation']=station['token']
            teacher.patch('/api/classroom/stations/'+station['id'],json={'approved':True},headers=HEAD)
        login(student,'student');sid=student.get('/api/me').json()['id']
        assert post(student,'/classroom/presence',{'sharing':False}).status_code==200
        post(teacher,'/classroom/presence',{'sharing':True})
        assert post(teacher,'/classroom/links',{'students':[sid],'mode':'audio'}).status_code==200
        assert post(teacher,'/classroom/links',{'students':[sid],'mode':'demo'}).status_code==200
        assert post(teacher,'/classroom/links',{'students':[sid],'mode':'help'}).status_code==400

def test_miniquestions_shared_with_teacher():
    with TestClient(app) as c:
        login(c,'admin')
        admin_questions=c.get('/api/learning/questions').json()
        assert admin_questions
        post(c,'/logout');login(c,'teacher')
        assert {q['id'] for q in c.get('/api/learning/questions').json()}=={q['id'] for q in admin_questions}
        q=admin_questions[0]
        r=c.put('/api/learning/questions/'+str(q['id']),json=q,headers=HEAD)
        assert r.status_code==200,r.text

def test_demo_students_have_different_playable_levels(monkeypatch):
    monkeypatch.setenv('DEMO_CLASSROOM','1')
    with TestClient(app) as c:
        for n in [1,3,5]:
            r=post(c,'/login',{'login':f'student{n}','password':'12345678'});assert r.status_code==200,r.text
            info=c.get('/api/learning/profile').json()
            assert info['intro_complete']==(n!=1)
            assert info['recommended_stars']==n
            assignments=c.get('/api/assignments').json()
            selected=next(a for a in assignments if a['title'].startswith('★'*n+' ·'))
            p=post(c,f'/assignments/{selected["id"]}/start').json()
            assert p['task']['limit_seconds']=={1:90,3:45,5:30}[n]
            assert bool(p['card']['house'])==(n<4)
            mini=post(c,f'/learning/drill/{p["id"]}').json()
            assert mini and mini['level']=={1:0,3:2,5:5}[n]
            post(c,'/logout')
        login(c,'teacher')
        rows=c.get('/api/teacher/students').json()
        assert any(r.get('demo') and r['score']==97 for r in rows)
