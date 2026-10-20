"""Read the supplied matrix without replacing the existing incident catalogue."""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl


def extract(source):
    sheet = openpyxl.load_workbook(source, data_only=True)['Лист1']
    columns = []
    for col in range(15, 91):
        labels = []
        for row in (1, 2, 3):
            cell = sheet.cell(row, col)
            for merged in sheet.merged_cells.ranges:
                if cell.coordinate in merged:
                    cell = sheet.cell(merged.min_row, merged.min_col)
                    break
            if cell.value and str(cell.value).strip() not in labels:
                labels.append(str(cell.value).strip())
        columns.append({'column': openpyxl.utils.get_column_letter(col), 'headers': labels})
    rows = []
    for row in sheet.iter_rows(min_row=4):
        values = [c.value for c in row]
        if not isinstance(values[0], (int, float)) or not isinstance(values[10], str):
            continue
        code = str(int(sum((values[i] or 0)*factor for i, factor in enumerate([1000000, 10000, 100, 1]))))
        cells = {}
        for c in row[14:90]:
            if c.value is not None or c.comment:
                cells[c.column_letter] = {'value': c.value, **({'note': c.comment.text} if c.comment else {})}
        rows.append({'code': code, 'title': values[10].strip(), 'row': row[0].row,
                     'traits': [str(v).strip() for v in values[6:10] if v is not None],
                     'scenario': values[12], 'lead_service': values[13], 'cells': cells})
    return {'version': 1, 'source': {'file': source.name, 'sheet': sheet.title,
            'sha256': hashlib.sha256(source.read_bytes()).hexdigest()},
            'columns': columns, 'rows': rows}


if __name__ == '__main__':
    import argparse
    parser=argparse.ArgumentParser(description='Extract the supplied classifier routing matrix without changing the workbook.')
    parser.add_argument('source',type=Path,help='Source XLSX workbook')
    parser.add_argument('target',type=Path,help='New JSON output file')
    args=parser.parse_args()
    source,target=args.source,args.target
    if target.exists():
        raise SystemExit('Choose a new output path; the reference file is never overwritten.')
    target.write_text(json.dumps(extract(source), ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
